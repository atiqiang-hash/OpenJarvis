# -*- coding: utf-8 -*-
"""Generate the Word usage tutorial for the Serenity dashboard."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---- CJK-friendly default font ----
def set_cjk(style, font="Microsoft YaHei", size=10.5):
    style.font.name = font
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font)
    rfonts.set(qn('w:ascii'), font)
    rfonts.set(qn('w:hAnsi'), font)

set_cjk(doc.styles['Normal'])
for s in ['Heading 1', 'Heading 2', 'Heading 3', 'Title', 'List Bullet', 'List Number']:
    try:
        set_cjk(doc.styles[s], size=doc.styles[s].font.size.pt if doc.styles[s].font.size else 13)
    except Exception:
        pass

GREEN = RGBColor(0x10, 0x9a, 0x66)
DIM = RGBColor(0x60, 0x60, 0x60)

def h(text, level=1):
    p = doc.add_heading(text, level=level)
    return p

def para(text="", bold=False, italic=False, color=None, size=10.5, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

def numbered(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Number')
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

# ============ COVER ============
t = doc.add_heading('Serenity「卡点猎手」股市分析仪表盘', level=0)
para('—— 使用教程 / User Guide', bold=True, color=GREEN, size=13)
para('基于 X 账号 @aleabitoreddit（Serenity）公开市场分析方法论的研究性可视化工具', color=DIM, size=10)
para('版本 v1.2 · 已接入 Finnhub 实时美股行情 · 生成日期 2026-06-16', color=DIM, size=9)
doc.add_paragraph()

# ============ 1 ============
h('一、这是什么', 1)
para('这是一个单文件的网页版股市分析仪表盘（index.html），把 Serenity 分析市场的「底层逻辑」做成了一套可交互的分析框架。核心思路是：')
para('不去追 NVIDIA 这类显性龙头，而是把超大厂（Mag7）的万亿级 AI 资本开支，沿供应链逆向拆解，'
     '找到那个「大厂愿意付任何价钱也要保供」的单点瓶颈（chokepoint），在机构轮动到达之前，'
     '布局被严重错配的微 / 小市值标的。', italic=True)
para('仪表盘包含六大模块：底层逻辑对照、卡点供应链七层图、实时标的看板、14 条核心原则、'
     '选股评分器、以及风险提示。')

# ============ 2 ============
h('二、文件清单', 1)
bullet('—— 仪表盘主程序，双击即可用浏览器打开（这就是全部，单文件、零安装）。', bold_prefix='index.html ')
bullet('—— 本使用教程（你正在看的文档）。', bold_prefix='使用教程.docx ')
bullet('—— 纯文本快速说明（备用）。', bold_prefix='README.txt ')
para('整个项目就是一个文件夹，复制到任何地方（包括移动 SSD）都能直接用，不依赖网络安装、不需要服务器。', color=DIM, size=9.5)

# ============ 3 ============
h('三、如何打开', 1)
numbered('找到文件夹里的 index.html。')
numbered('双击它 —— 系统会用默认浏览器打开（推荐 Chrome / Edge / Safari 等现代浏览器）。')
numbered('打开后页面会自动向 Finnhub 拉取一次实时行情；看到顶部状态灯变绿、显示「已更新 X/20」即表示成功。')
para('说明：仪表盘本身可完全离线浏览（框架、图表、评分器都不需要网络）；'
     '只有「实时行情」这一项需要联网访问 Finnhub。', color=DIM, size=9.5)

# ============ 4 ============
h('四、界面分区导览', 1)

h('1. 底层逻辑', 2)
para('左右对照「传统 / 共识打法」与「卡点猎手打法」，一眼看懂这套方法和大众做法的根本区别。')

h('2. 卡点供应链（七层）', 2)
para('把 AI 供应链从「超大厂资本开支」一路拆到「衬底 / 原料」，分成 7 个卡点层。'
     '点击任意一层，下方会展开该层的逻辑说明与代表标的。越往上游，市值越小、错配通常越大。')
# layers table
layers = [
    ('💡 光互连 / CPO 光子学', 'AXTI · LITE · COHR · AAOI · POET · AEHR'),
    ('🔗 连接 / 重定时 / 时钟', 'CRDO · ALAB · SITM · MRVL'),
    ('🧪 化合物衬底 / 原料', 'AXTI · MTSI · COHR'),
    ('🧠 存储 / HBM / NAND', 'MU · SNDK · INTC'),
    ('⚡ 电力 / 电网 / 散热', 'VRT · GEV · POWL'),
    ('🤖 物理 AI / 机器人 / 航天', 'RKLB · AMD · ARM'),
    ('☁️ Neocloud / 算力融资', 'NBIS · CRWV · CRCL'),
]
tb = doc.add_table(rows=1, cols=2)
tb.style = 'Light Grid Accent 1'
tb.rows[0].cells[0].paragraphs[0].add_run('卡点层').bold = True
tb.rows[0].cells[1].paragraphs[0].add_run('代表标的').bold = True
for name, tks in layers:
    c = tb.add_row().cells
    c[0].text = name; c[1].text = tks

h('3. 标的看板（实时）', 2)
para('20 只标的的实时行情看板。可按卡点层筛选、点表头排序。各列含义如下：')
cols = [
    ('标的', '代码 + 公司名 + 时间维度'),
    ('档', '信念档：S（最高）/ A / B，按错配确定性与赔率分级'),
    ('卡点层', '所属的供应链卡点层'),
    ('现价·实时', '来自 Finnhub 的实时美股价格'),
    ('日涨跌', '当日涨跌幅，绿涨红跌（▲ / ▼）'),
    ('日内区间', '小色条标记现价在「当日最低 ↔ 最高」之间的位置'),
    ('参考 → 目标', 'Serenity 公开发帖里提到的参考价 → 目标价（静态，可能已过时）'),
    ('距目标·实时', '用实时现价算出的「到目标价还有多少空间」'),
    ('核心论点', '该标的的卡点逻辑一句话'),
]
tb2 = doc.add_table(rows=1, cols=2)
tb2.style = 'Light Grid Accent 1'
tb2.rows[0].cells[0].paragraphs[0].add_run('列名').bold = True
tb2.rows[0].cells[1].paragraphs[0].add_run('含义').bold = True
for n, m in cols:
    c = tb2.add_row().cells
    c[0].text = n; c[1].text = m

h('4. 14 条核心原则', 2)
para('把这套打法拆成 14 个可复用、可迁移的判断模块，例如「瓶颈猎杀」「多跳 BOM 测绘」'
     '「签约 ARR vs 市值错配」「稀释 / ATM 日历 = 否决项」「机构滞后窗口」等。')

h('5. Serenity 选股评分器', 2)
para('一个交互式自查工具。逐项勾选你正在研究的标的是否满足 14 条检查清单，'
     '右侧实时给出「卡点契合度」评分与结论：')
bullet('0 项 — 等待评估')
bullet('< 45% — 回避（红）：卡点特征不足')
bullet('45%–70% — 观察 / 待催化（黄）')
bullet('70%–86% — 符合 / 定义风险建仓（青）')
bullet('> 86% — 高信念 / 非对称机会（紫）')
para('该评分仅为框架自查工具，不构成任何投资建议。', italic=True, color=DIM, size=9.5)

h('6. 反模式与风险', 2)
para('列出 Serenity 主动回避的 6 类信号（纯技术面、头衔式评论、把内部人卖出当信号、'
     '混淆供应链层级、追情绪、事件后抢跑），以及完整的风险免责声明。')

# ============ 5 ============
h('五、实时行情（Finnhub）使用说明', 1)
para('标的看板顶部有一条「实时控制条」，从左到右：')
bullet('—— 灰=就绪、黄=加载中 / 限流、绿=成功、红=失败。', bold_prefix='状态灯 ')
bullet('—— 显示「已更新 X/20 · 组合日内均值 ±x% · 时间戳」。', bold_prefix='状态文字 ')
bullet('—— 手动拉取一次最新实时报价。', bold_prefix='↻ 刷新行情 ')
bullet('—— 勾选后每 60 秒自动刷新一次（避开免费层 60 次/分钟的限流）。', bold_prefix='自动 60s ')
bullet('—— 替换 / 自定义你自己的 Finnhub API Key（见下一节）。', bold_prefix='🔑 API Key ')
para('数据源为 Finnhub 免费层（美股、60 calls/min）。页面已内置一个免费 Key，开箱即用。', color=DIM, size=9.5)

# ============ 6 ============
h('六、替换 / 自定义 API Key', 1)
numbered('打开 https://finnhub.io/dashboard 免费注册，复制你自己的 API Key。')
numbered('在仪表盘点「🔑 API Key」按钮，把 Key 粘进去、确定。')
numbered('Key 会保存在你这台电脑的浏览器本地（localStorage），不会写回文件、也不会上传。')
numbered('想恢复内置默认 Key，再点一次「🔑 API Key」，把输入框清空、确定即可。')
para('提示：内置 Key 是免费、只读、限流的公共额度。如果多人同时用导致限流（状态灯变黄），'
     '换成你自己的 Key 会更稳定。', color=DIM, size=9.5)

# ============ 7 ============
h('七、常见问题 / 排错', 1)
bullet('点「刷新行情」即可；若仍空，多为该时段限流或浏览器拦截了跨域请求，稍等再试。',
       bold_prefix='Q：某只标的一直显示「—」？ ')
bullet('免费层限流（60 次/分钟）。关掉「自动 60s」或换自己的 Key。',
       bold_prefix='Q：状态灯变黄、提示限流？ ')
bullet('检查电脑是否联网；公司 / 学校网络可能屏蔽 finnhub.io；换网络或换 Key 重试。',
       bold_prefix='Q：状态灯变红、拉取失败？ ')
bullet('正常。框架、图表、评分器都能离线用，只有实时行情需要联网。',
       bold_prefix='Q：没网能用吗？ ')
bullet('目标价 / 参考价是 Serenity 历史发帖里的数字，是静态的、可能已过时；只有「现价 / 日涨跌 / 距目标」是实时的。',
       bold_prefix='Q：目标价为什么不变？ ')

# ============ 8 ============
h('八、保存到移动 SSD 的步骤', 1)
numbered('把收到的压缩包（Serenity-卡点仪表盘.zip）下载到电脑。')
numbered('把移动 SSD 插上电脑。')
numbered('解压压缩包，得到「Serenity-卡点仪表盘」文件夹。')
numbered('把整个文件夹拖进 SSD 里你想放的位置即可。')
numbered('以后在 SSD 里双击 index.html 就能用（联网时自动出实时行情）。')

# ============ 9 ============
h('九、重要风险与免责声明', 1)
para('请务必阅读：', bold=True)
numbered('本工具是对公开账号 @aleabitoreddit（Serenity）市场分析方法论的二次整理与可视化，'
         '并非其本人或任何机构发布，不构成投资、财务或交易建议。')
numbered('其自报收益（如 +3,600% / +4,502%）未经审计、含杠杆、波动极大，且存在严重的幸存者 / 选择偏差；'
         '他不管理基金、不披露 13F，所点名标的多为高波动微 / 小市值，可单日暴涨亦可腰斩（如 $BKKT 曾从 $40 跌至 $8）。')
numbered('现价 / 日涨跌为 Finnhub 实时报价；参考价 / 目标价取自其历史发帖、可能已过时。'
         '任何决策请以你自己的尽调与实时数据为准。')
numbered('粉丝放大效应可能造成短期情绪冲高，存在在高位接盘的风险。投资有风险，入市需谨慎。')

# ============ 10 ============
h('十、数据来源', 1)
bullet('X 账号：x.com/aleabitoreddit')
bullet('方法论存档：github.com/yan-labs/serenity-aleabitoreddit')
bullet('14 原则整理：bearsavings.com/blog/who-is-serenity-aleabitoreddit')
bullet('实时行情：finnhub.io')

doc.add_paragraph()
para('—— 教程结束。祝研究愉快，理性投资。', color=GREEN, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

out = '/home/user/OpenJarvis/Serenity-卡点仪表盘/使用教程.docx'
doc.save(out)
print('saved', out)
